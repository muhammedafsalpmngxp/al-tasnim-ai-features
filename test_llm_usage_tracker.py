"""Tests for llm_usage_tracker.py.

Kept at the project root, next to the module it tests, rather than inside
``backend/tests/`` -- that keeps this test file discoverable and runnable
completely independently of the backend's own test suite (``pytest.ini``
there only looks under ``backend/tests``), matching the tracker's own
independence from the rest of the project. Run directly with:

    python -m pytest test_llm_usage_tracker.py -v

Every test points the module at a throwaway ``.env``/log file under
``tmp_path`` (via ``monkeypatch``), so no test ever touches the real
``backend/.env`` or the real ``llm_usage_log.xlsx`` at the project root.
"""

from __future__ import annotations

from datetime import datetime, timezone

import openpyxl
import pytest

import llm_usage_tracker as tracker


@pytest.fixture(autouse=True)
def _isolated_files(tmp_path, monkeypatch):
    """Every test gets its own env file and log file; nothing here ever
    touches the real project-root llm_usage_log.xlsx or backend/.env."""
    env_file = tmp_path / "fake.env"
    env_file.write_text("LLM_PROVIDER=openai\nLLM_MODEL=gpt-4o-mini\n", encoding="utf-8")
    log_file = tmp_path / "usage.xlsx"
    monkeypatch.setattr(tracker, "_ENV_FILE", env_file)
    monkeypatch.setattr(tracker, "_LOG_FILE", log_file)
    return env_file, log_file


<<<<<<< HEAD
=======
def _column(name):
    """The 0-based index of a column, by its header name.

    Indexed by NAME rather than by position on purpose: this file used to
    assert on fixed offsets, and adding the routing columns (System, Node,
    Capability) broke seven tests that were not testing anything about
    routing. A test that fails when an unrelated column is added is a test
    that has to be rewritten every time the log grows.
    """
    return tracker._HEADERS.index(name)


def _cell(row, name):
    return row[_column(name)]


def _letter(name):
    from openpyxl.utils import get_column_letter

    return get_column_letter(_column(name) + 1)


>>>>>>> 8e557ed (Update Daily Report Agentic implementation)
def _rows(log_file):
    """Every row's cell values, each as a list (so it compares equal to
    ``tracker._HEADERS``, itself a plain list)."""
    workbook = openpyxl.load_workbook(log_file)
    sheet = workbook.active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


class TestEnvParsing:
    def test_reads_provider_and_model_from_the_env_file(self, _isolated_files):
        provider, model = tracker._current_provider_and_model()
        assert provider == "openai"
        assert model == "gpt-4o-mini"

    def test_missing_env_file_reports_unknown_rather_than_raising(self, tmp_path, monkeypatch):
        monkeypatch.setattr(tracker, "_ENV_FILE", tmp_path / "does-not-exist.env")
        provider, model = tracker._current_provider_and_model()
        assert provider == "unknown"
        assert model == "unknown"

    def test_quoted_and_commented_values_are_handled(self, tmp_path, monkeypatch):
        env_file = tmp_path / "fake.env"
        env_file.write_text(
            "# a comment\nLLM_PROVIDER=\"groq\"\n\nLLM_MODEL='openai/gpt-oss-120b'\n",
            encoding="utf-8",
        )
        monkeypatch.setattr(tracker, "_ENV_FILE", env_file)
        provider, model = tracker._current_provider_and_model()
        assert provider == "groq"
        assert model == "openai/gpt-oss-120b"


class TestLoggingOneCall:
    def test_a_first_call_creates_the_file_with_headers_data_and_totals(self, _isolated_files):
        _, log_file = _isolated_files
        ok = tracker.log_usage(input_tokens=100, output_tokens=20, duration_seconds=1.5)
        assert ok is True

        rows = _rows(log_file)
        assert rows[0] == tracker._HEADERS
<<<<<<< HEAD
        assert rows[1][1:] == ["openai", "gpt-4o-mini", 1.5, 100, 20, None]
        totals = rows[2]
        assert totals[0] == '="TOTAL calls: "&COUNTA(A2:A2)'
        assert totals[3] == "=SUM(D2:D2)"
        assert totals[4] == "=SUM(E2:E2)"
        assert totals[5] == "=SUM(F2:F2)"
        assert totals[6] == "=SUM(G2:G2)"
=======
        row = rows[1]
        assert _cell(row, "Provider") == "openai"
        assert _cell(row, "Model") == "gpt-4o-mini"
        assert _cell(row, "Duration (s)") == 1.5
        assert _cell(row, "Input Tokens") == 100
        assert _cell(row, "Output Tokens") == 20
        assert _cell(row, "Input Cache Tokens") is None
        # A call made without saying which model system it belongs to is the
        # explanation layer -- every caller that predates the compiler is.
        assert _cell(row, "System") == tracker.SYSTEM_FAST
        assert _cell(row, "Node") == tracker.NODE_FAST_EXPLANATION
        # Blank, because an explanation is about a request rather than about a
        # capability -- not a capability named "".
        assert _cell(row, "Capability") is None
        # One subtotal row for the one model, then the grand total.
        assert len(rows) == 4
        assert rows[2][0] == '="TOTAL gpt-4o-mini calls: "&COUNTIF($F$2:$F$2,"gpt-4o-mini")'
        totals = rows[3]
        assert totals[0] == '="TOTAL calls: "&COUNTA(A2:A2)'
        for name in ("Duration (s)", "Input Tokens", "Output Tokens", "Input Cache Tokens"):
            letter = _letter(name)
            assert _cell(totals, name) == f"=SUM({letter}2:{letter}2)"
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)

    def test_input_cache_tokens_is_recorded_when_given(self, _isolated_files):
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=100, output_tokens=20, input_cache_tokens=64, duration_seconds=1.5)
        rows = _rows(log_file)
        assert rows[1][-1] == 64

    def test_missing_token_counts_are_stored_as_blank_not_zero(self, _isolated_files):
        """A provider that doesn't report usage must not look like it used
        zero tokens -- None and 0 mean different things here."""
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=None, output_tokens=None, duration_seconds=0.9)
        rows = _rows(log_file)
<<<<<<< HEAD
        assert rows[1][4] is None
        assert rows[1][5] is None
=======
        assert _cell(rows[1], "Input Tokens") is None
        assert _cell(rows[1], "Output Tokens") is None
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)


class TestTotalsRowStaysAtTheBottom:
    def test_three_calls_keep_the_totals_row_last_with_a_growing_range(self, _isolated_files):
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=10, output_tokens=1, duration_seconds=0.1)
        tracker.log_usage(input_tokens=20, output_tokens=2, duration_seconds=0.2)
        tracker.log_usage(input_tokens=30, output_tokens=3, duration_seconds=0.3)

        rows = _rows(log_file)
<<<<<<< HEAD
        assert len(rows) == 5  # header + 3 data rows + 1 totals row
        assert rows[1][4] == 10 and rows[2][4] == 20 and rows[3][4] == 30
        totals = rows[4]
        assert totals[0] == '="TOTAL calls: "&COUNTA(A2:A4)'
        assert totals[3] == "=SUM(D2:D4)"
        assert totals[4] == "=SUM(E2:E4)"
=======
        assert len(rows) == 6  # header + 3 data rows + 1 model subtotal + grand total
        assert [_cell(r, "Input Tokens") for r in rows[1:4]] == [10, 20, 30]
        totals = rows[5]
        assert totals[0] == '="TOTAL calls: "&COUNTA(A2:A4)'
        for name in ("Duration (s)", "Input Tokens"):
            letter = _letter(name)
            assert _cell(totals, name) == f"=SUM({letter}2:{letter}4)"
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)

    def test_the_actual_sums_are_arithmetically_correct(self, _isolated_files):
        """Formulas aren't evaluated by openpyxl -- recompute by hand from
        the raw cell values to prove the *ranges* the formulas reference
        actually cover the right data, not just that a formula exists."""
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=10, output_tokens=1, input_cache_tokens=5, duration_seconds=0.1)
        tracker.log_usage(input_tokens=20, output_tokens=2, input_cache_tokens=None, duration_seconds=0.2)
        tracker.log_usage(input_tokens=30, output_tokens=3, input_cache_tokens=7, duration_seconds=0.3)

        rows = _rows(log_file)
        data_rows = rows[1:4]
<<<<<<< HEAD
        assert sum(r[4] for r in data_rows) == 60  # Input Tokens
        assert sum(r[5] for r in data_rows) == 6  # Output Tokens
        assert sum(r[6] for r in data_rows if r[6] is not None) == 12  # Input Cache Tokens
        assert round(sum(r[3] for r in data_rows), 6) == 0.6  # Duration
=======
        assert sum(_cell(r, "Input Tokens") for r in data_rows) == 60
        assert sum(_cell(r, "Output Tokens") for r in data_rows) == 6
        assert (
            sum(
                _cell(r, "Input Cache Tokens")
                for r in data_rows
                if _cell(r, "Input Cache Tokens") is not None
            )
            == 12
        )
        assert round(sum(_cell(r, "Duration (s)") for r in data_rows), 6) == 0.6
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)


class TestMigrationFromAnOlderFile:
    """Simulates upgrading a log file written before this version existed:
    no "Input Cache Tokens" column, no totals row."""

    def _write_old_format_file(self, log_file):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Timestamp", "Provider", "Model", "Duration (s)", "Input Tokens", "Output Tokens"])
        sheet.append(["2026-08-01T00:00:00+00:00", "groq", "openai/gpt-oss-120b", 2.0, 500, 80])
        workbook.save(log_file)

    def test_existing_rows_are_preserved_with_a_blank_new_column(self, _isolated_files):
        _, log_file = _isolated_files
        self._write_old_format_file(log_file)

        tracker.log_usage(input_tokens=100, output_tokens=10, input_cache_tokens=5, duration_seconds=1.0)

        rows = _rows(log_file)
        assert rows[0] == tracker._HEADERS
        old_row = rows[1]
<<<<<<< HEAD
        assert old_row[:6] == ["2026-08-01T00:00:00+00:00", "groq", "openai/gpt-oss-120b", 2.0, 500, 80]
        assert old_row[6] is None  # never captured for this historical row
        new_row = rows[2]
        assert new_row[4:] == [100, 10, 5]
=======
        assert old_row[0] == "2026-08-01T00:00:00+00:00"
        assert _cell(old_row, "Provider") == "groq"
        assert _cell(old_row, "Model") == "openai/gpt-oss-120b"
        assert _cell(old_row, "Duration (s)") == 2.0
        assert _cell(old_row, "Input Tokens") == 500
        assert _cell(old_row, "Output Tokens") == 80
        # Columns this version added were never captured for a historical row.
        # They are left blank rather than guessed at: a call whose routing
        # nobody recorded is not a call that was routed to the default.
        for added in ("Input Cache Tokens", "System", "Node", "Capability"):
            assert _cell(old_row, added) is None
        new_row = rows[2]
        assert _cell(new_row, "Input Tokens") == 100
        assert _cell(new_row, "Output Tokens") == 10
        assert _cell(new_row, "Input Cache Tokens") == 5
        assert _cell(new_row, "System") == tracker.SYSTEM_FAST
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)

    def test_a_totals_row_is_added_covering_old_and_new_rows_together(self, _isolated_files):
        _, log_file = _isolated_files
        self._write_old_format_file(log_file)
        tracker.log_usage(input_tokens=100, output_tokens=10, duration_seconds=1.0)

        rows = _rows(log_file)
        totals = rows[-1]
        assert totals[0] == '="TOTAL calls: "&COUNTA(A2:A3)'
<<<<<<< HEAD
        assert totals[4] == "=SUM(E2:E3)"
=======
        letter = _letter("Input Tokens")
        assert _cell(totals, "Input Tokens") == f"=SUM({letter}2:{letter}3)"
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)

    def test_re_running_against_an_already_current_file_changes_nothing_but_the_new_row(
        self, _isolated_files
    ):
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=1, output_tokens=1, duration_seconds=0.1)
        tracker.log_usage(input_tokens=2, output_tokens=2, duration_seconds=0.2)
        rows = _rows(log_file)
        assert rows[0] == tracker._HEADERS
<<<<<<< HEAD
        assert len(rows) == 4  # header + 2 data + totals -- no duplicate migration
=======
        assert len(rows) == 5  # header + 2 data + model subtotal + grand total -- no duplicate migration
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)


class TestNeverRaises:
    def test_a_file_that_cannot_be_written_returns_false_not_an_exception(self, tmp_path, monkeypatch):
        # A directory where the log file path should be forces a write error.
        blocked = tmp_path / "blocked"
        blocked.mkdir()
        monkeypatch.setattr(tracker, "_LOG_FILE", blocked)  # a directory, not a file
        result = tracker.log_usage(input_tokens=1, output_tokens=1, duration_seconds=0.1)
        assert result is False


# ---------------------------------------------------------------------------
# The authoritative sync
#
# No admin key and no network here: `_get_json` is replaced with recorded
# response shapes from OpenAI's documented usage/cost endpoints, so what is
# asserted is how this script reads them and what it writes.
# ---------------------------------------------------------------------------


def _sheet_rows(log_file, title):
    workbook = openpyxl.load_workbook(log_file)
    if title not in workbook.sheetnames:
        return None
    return [[cell.value for cell in row] for row in workbook[title].iter_rows()]


def _usage_page(day_start, results, *, next_page=None):
    """One page of the completions-usage endpoint, as documented."""
    return {
        "object": "page",
        "data": [
            {
                "object": "bucket",
                "start_time": day_start,
                "end_time": day_start + 86400,
                "results": results,
            }
        ],
        "has_more": next_page is not None,
        "next_page": next_page,
    }


def _usage_result(model, *, requests=1, input_tokens=0, output_tokens=0, cached=0):
    return {
        "object": "organization.usage.completions.result",
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "input_cached_tokens": cached,
        "num_model_requests": requests,
        "project_id": None,
        "model": model,
    }


#: Bucket starts, as the API reports them: Unix seconds at UTC midnight.
#: Derived rather than typed as a magic number, so the dates asserted below
#: and the timestamps fed in can never drift apart.
DAY_ONE_DATE = "2026-09-16"
DAY_TWO_DATE = "2026-09-17"
DAY_ONE = int(datetime(2026, 9, 16, tzinfo=timezone.utc).timestamp())
DAY_TWO = DAY_ONE + 86400


@pytest.fixture
def admin_env(_isolated_files, monkeypatch):
    env_file, log_file = _isolated_files
    env_file.write_text(
        "LLM_PROVIDER=openai\nLLM_MODEL=gpt-4o-mini\nOPENAI_ADMIN_KEY=sk-admin-test\n",
        encoding="utf-8",
    )
    return env_file, log_file


@pytest.fixture
def fake_api(monkeypatch):
    """Serves recorded responses and records the URLs asked for."""
    calls = {"urls": [], "usage": [], "costs": []}

    def responder(url, admin_key):
        calls["urls"].append(url)
        assert admin_key == "sk-admin-test"
        if "/organization/costs" in url:
            queue = calls["costs"]
        else:
            queue = calls["usage"]
        if not queue:
            return {"object": "page", "data": [], "has_more": False, "next_page": None}
        return queue.pop(0)

    monkeypatch.setattr(tracker, "_get_json", responder)
    return calls


class TestSyncConfiguration:
    def test_the_admin_key_is_read_from_the_env_file(self, admin_env):
        settings = tracker._usage_api_settings()
        assert settings["admin_key"] == "sk-admin-test"
        assert settings["base_url"] == tracker._DEFAULT_USAGE_BASE_URL
        assert settings["days"] == tracker._DEFAULT_SYNC_DAYS
        assert settings["project_ids"] == []

    def test_the_window_project_and_endpoint_are_configurable(self, _isolated_files):
        env_file, _ = _isolated_files
        env_file.write_text(
            "OPENAI_ADMIN_KEY=sk-admin-test\n"
            "OPENAI_USAGE_DAYS=7\n"
            "OPENAI_USAGE_PROJECT_IDS=proj_a, proj_b\n"
            "OPENAI_USAGE_BASE_URL=https://example.invalid/v1\n",
            encoding="utf-8",
        )
        settings = tracker._usage_api_settings()
        assert settings["days"] == 7
        assert settings["project_ids"] == ["proj_a", "proj_b"]
        assert settings["base_url"] == "https://example.invalid/v1"

    def test_a_nonsense_window_falls_back_to_the_default(self, _isolated_files):
        env_file, _ = _isolated_files
        env_file.write_text("OPENAI_ADMIN_KEY=k\nOPENAI_USAGE_DAYS=not-a-number\n", encoding="utf-8")
        assert tracker._usage_api_settings()["days"] == tracker._DEFAULT_SYNC_DAYS

    def test_without_an_admin_key_nothing_is_fetched_or_written(self, _isolated_files, monkeypatch):
        _, log_file = _isolated_files

        def never_called(*args, **kwargs):  # pragma: no cover - must not run
            raise AssertionError("the API must not be called without an admin key")

        monkeypatch.setattr(tracker, "_get_json", never_called)
        with pytest.raises(tracker.UsageSyncError) as exc:
            tracker.sync_openai_usage()
        assert "OPENAI_ADMIN_KEY" in str(exc.value)
        assert not log_file.exists()

    def test_the_admin_key_is_never_echoed_in_an_error(self, admin_env, monkeypatch):
        def refused(url, admin_key):
            raise tracker.urllib.error.HTTPError(url, 401, "Unauthorized", {}, None)

        monkeypatch.setattr(tracker, "_get_json", tracker._get_json)
        monkeypatch.setattr(tracker.urllib.request, "urlopen", lambda *a, **k: refused("u", "k"))
        with pytest.raises(tracker.UsageSyncError) as exc:
            tracker.sync_openai_usage()
        message = str(exc.value)
        assert "sk-admin-test" not in message
        assert "ADMIN key" in message


class TestReadingWhatOpenAIReports:
    def test_usage_is_shaped_per_day_and_model(self, admin_env, fake_api):
        fake_api["usage"].append(
            _usage_page(
                DAY_ONE,
                [
                    _usage_result("gpt-4o-mini", requests=12, input_tokens=9000, output_tokens=1200, cached=512),
                    _usage_result("gpt-4o", requests=2, input_tokens=400, output_tokens=90),
                ],
            )
        )
        rows = tracker.fetch_openai_usage(tracker._usage_api_settings(), days=7)
        assert [row["model"] for row in rows] == ["gpt-4o", "gpt-4o-mini"]  # sorted
        mini = [row for row in rows if row["model"] == "gpt-4o-mini"][0]
        assert (mini["requests"], mini["input_tokens"], mini["output_tokens"]) == (12, 9000, 1200)
        assert mini["input_cached_tokens"] == 512
        assert mini["date"] == DAY_ONE_DATE

    def test_a_bucket_with_no_usage_is_not_written_as_a_row_of_zeroes(self, admin_env, fake_api):
        fake_api["usage"].append(_usage_page(DAY_ONE, [_usage_result("gpt-4o-mini", requests=0)]))
        assert tracker.fetch_openai_usage(tracker._usage_api_settings(), days=7) == []

    def test_every_page_is_followed_so_a_long_window_is_never_short(self, admin_env, fake_api):
        fake_api["usage"].append(
            _usage_page(DAY_ONE, [_usage_result("gpt-4o-mini", requests=1, input_tokens=10)], next_page="page-2")
        )
        fake_api["usage"].append(
            _usage_page(DAY_TWO, [_usage_result("gpt-4o-mini", requests=3, input_tokens=30)])
        )
        rows = tracker.fetch_openai_usage(tracker._usage_api_settings(), days=7)
        assert [row["date"] for row in rows] == [DAY_ONE_DATE, DAY_TWO_DATE]
        assert "page=page-2" in fake_api["urls"][1]

    def test_the_request_asks_for_daily_buckets_grouped_by_model(self, admin_env, fake_api):
        tracker.fetch_openai_usage(tracker._usage_api_settings(), days=7)
        url = fake_api["urls"][0]
        assert "/organization/usage/completions?" in url
        assert "bucket_width=1d" in url
        assert "group_by=model" in url
        assert "start_time=" in url

    def test_costs_are_summed_per_day(self, admin_env, fake_api):
        fake_api["costs"].append(
            {
                "object": "page",
                "data": [
                    {
                        "object": "bucket",
                        "start_time": DAY_ONE,
                        "end_time": DAY_TWO,
                        "results": [
                            {"amount": {"value": 0.12, "currency": "usd"}, "line_item": "gpt-4o-mini, input"},
                            {"amount": {"value": 0.03, "currency": "usd"}, "line_item": "gpt-4o-mini, output"},
                        ],
                    }
                ],
                "has_more": False,
                "next_page": None,
            }
        )
        costs = tracker.fetch_openai_costs(tracker._usage_api_settings(), days=7)
        assert costs == {DAY_ONE_DATE: pytest.approx(0.15)}


class TestSyncWritesTheWorkbook:
    @pytest.fixture
    def synced(self, admin_env, fake_api):
        _, log_file = admin_env
        # Two calls this app logged itself, on the same day OpenAI reports.
        when = datetime(2026, 9, 16, 9, 0, tzinfo=timezone.utc)
        tracker.log_usage(input_tokens=1000, output_tokens=100, duration_seconds=1.0, timestamp=when)
        tracker.log_usage(input_tokens=2000, output_tokens=200, duration_seconds=2.0, timestamp=when)
        fake_api["usage"].append(
            _usage_page(
                DAY_ONE,
                [_usage_result("gpt-4o-mini", requests=9, input_tokens=9000, output_tokens=900, cached=256)],
            )
        )
        fake_api["costs"].append(
            {
                "object": "page",
                "data": [
                    {
                        "object": "bucket",
                        "start_time": DAY_ONE,
                        "end_time": DAY_TWO,
                        "results": [{"amount": {"value": 0.25, "currency": "usd"}}],
                    }
                ],
                "has_more": False,
                "next_page": None,
            }
        )
        result = tracker.sync_openai_usage(days=7)
        return log_file, result

    def test_the_provider_sheet_holds_what_openai_reported(self, synced):
        log_file, _ = synced
        rows = _sheet_rows(log_file, tracker._SHEET_PROVIDER)
        assert rows[0] == tracker._PROVIDER_HEADERS
        assert rows[1][:7] == [DAY_ONE_DATE, "gpt-4o-mini", 9, 9000, 256, 900, 9900]
        assert rows[2][0] == '="TOTAL rows: "&COUNTA(A2:A2)'
        assert rows[2][3] == "=SUM(D2:D2)"

    def test_the_reconciliation_shows_what_this_app_never_saw(self, synced):
        log_file, _ = synced
        rows = _sheet_rows(log_file, tracker._SHEET_RECONCILIATION)
        assert rows[0] == tracker._RECONCILIATION_HEADERS
        day = rows[1]
        assert day[0] == DAY_ONE_DATE
        assert day[1] == 9          # OpenAI billed nine requests
        assert day[2] == 2          # this app logged two calls
        assert day[3] == 7          # seven it never saw -- the whole point
        assert (day[4], day[5]) == (9000, 3000)   # input tokens, billed vs logged
        assert (day[6], day[7]) == (900, 300)     # output tokens, billed vs logged
        assert day[8] == pytest.approx(0.25)

    def test_the_per_call_sheet_is_left_exactly_as_it_was(self, synced):
        log_file, _ = synced
        rows = _sheet_rows(log_file, tracker._SHEET_CALLS)
        assert rows[0] == tracker._HEADERS
<<<<<<< HEAD
        assert len(rows) == 4  # header + the two logged calls + their totals row
=======
        assert len(rows) == 5  # header + the two logged calls + model subtotal + grand total
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)
        assert rows[-1][0] == '="TOTAL calls: "&COUNTA(A2:A3)'

    def test_the_sync_reports_the_gap_it_found(self, synced):
        _, result = synced
        assert result["provider_requests"] == 9
        assert result["logged_calls"] == 2
        assert result["cost_usd"] == pytest.approx(0.25)
        assert result["days"] == 7

    def test_a_day_openai_billed_but_this_app_never_logged_is_still_a_row(self, admin_env, fake_api):
        _, log_file = admin_env
        fake_api["usage"].append(
            _usage_page(DAY_ONE, [_usage_result("gpt-4o-mini", requests=4, input_tokens=400, output_tokens=40)])
        )
        tracker.sync_openai_usage(days=7)
        rows = _sheet_rows(log_file, tracker._SHEET_RECONCILIATION)
        assert rows[1][0] == DAY_ONE_DATE
        assert (rows[1][1], rows[1][2], rows[1][3]) == (4, 0, 4)

    def test_a_day_with_no_reported_cost_is_left_blank_not_zero(self, admin_env, fake_api):
        _, log_file = admin_env
        fake_api["usage"].append(
            _usage_page(DAY_ONE, [_usage_result("gpt-4o-mini", requests=1, input_tokens=10, output_tokens=1)])
        )
        tracker.sync_openai_usage(days=7)
        rows = _sheet_rows(log_file, tracker._SHEET_RECONCILIATION)
        assert rows[1][8] is None

    def test_syncing_the_same_window_again_replaces_rather_than_duplicates(self, admin_env, fake_api):
        _, log_file = admin_env
        for _ in range(2):
            fake_api["usage"].append(
                _usage_page(DAY_ONE, [_usage_result("gpt-4o-mini", requests=4, input_tokens=400, output_tokens=40)])
            )
            tracker.sync_openai_usage(days=7)
        rows = _sheet_rows(log_file, tracker._SHEET_PROVIDER)
        assert len([row for row in rows if row[0] == DAY_ONE_DATE]) == 1

    def test_a_call_logged_after_a_sync_still_lands_on_the_per_call_sheet(self, synced):
        # Regression: with three sheets in the workbook, "the active sheet" is
        # whichever one Excel last had selected -- a new row must go to the
        # per-call sheet by name, never to whatever that happens to be.
        log_file, _ = synced
        workbook = openpyxl.load_workbook(log_file)
        workbook.active = workbook.sheetnames.index(tracker._SHEET_RECONCILIATION)
        workbook.save(log_file)

        tracker.log_usage(input_tokens=7, output_tokens=7, duration_seconds=0.7)

        calls = _sheet_rows(log_file, tracker._SHEET_CALLS)
<<<<<<< HEAD
        assert calls[-2][4] == 7  # the new row, directly above the totals row
        assert len(_sheet_rows(log_file, tracker._SHEET_RECONCILIATION)[0]) == len(
            tracker._RECONCILIATION_HEADERS
        )
=======
        # the new row, directly above the model subtotal and the grand total
        assert _cell(calls[-3], "Input Tokens") == 7
        assert len(_sheet_rows(log_file, tracker._SHEET_RECONCILIATION)[0]) == len(
            tracker._RECONCILIATION_HEADERS
        )


class TestModelSystemRouting:
    """Two model systems now spend tokens through this log, and telling them
    apart is the question an operator actually asks: what is the SQL compiler
    costing me, as opposed to the report itself?"""

    def test_a_reasoning_call_is_recorded_under_its_own_system_and_node(self, _isolated_files):
        _, log_file = _isolated_files
        tracker.log_usage(
            input_tokens=9000,
            output_tokens=800,
            input_cache_tokens=4096,
            duration_seconds=12.0,
            system=tracker.SYSTEM_REASONING,
            node=tracker.NODE_REASONING_SQL_AUTHOR,
            capability="DAILY_DETAILS",
            provider="openai",
            model="a-reasoning-model",
        )
        row = _rows(log_file)[1]
        assert _cell(row, "System") == tracker.SYSTEM_REASONING
        assert _cell(row, "Node") == tracker.NODE_REASONING_SQL_AUTHOR
        assert _cell(row, "Capability") == "DAILY_DETAILS"
        # The compiler's model is NOT necessarily LLM_MODEL, so it passes its
        # own rather than letting the .env answer for it.
        assert _cell(row, "Model") == "a-reasoning-model"
        assert _cell(row, "Provider") == "openai"

    def test_the_two_systems_sit_side_by_side_and_can_be_totalled_apart(self, _isolated_files):
        _, log_file = _isolated_files
        tracker.log_usage(input_tokens=900, output_tokens=100, duration_seconds=1.0)
        tracker.log_usage(
            input_tokens=9000,
            output_tokens=900,
            duration_seconds=10.0,
            system=tracker.SYSTEM_REASONING,
            node=tracker.NODE_REASONING_VERIFIER,
            capability="WELL_ACTIVITY",
        )
        data = _rows(log_file)[1:3]
        by_system = {_cell(r, "System"): r for r in data}
        assert set(by_system) == {tracker.SYSTEM_FAST, tracker.SYSTEM_REASONING}
        assert _cell(by_system[tracker.SYSTEM_FAST], "Input Tokens") == 900
        assert _cell(by_system[tracker.SYSTEM_REASONING], "Input Tokens") == 9000

    def test_the_routing_is_never_inferred_from_the_provider(self, _isolated_files):
        # Two systems can be pointed at one provider, and one system's model
        # can change without anybody editing the tracker. The purpose is
        # recorded at the call site or not at all.
        _, log_file = _isolated_files
        tracker.log_usage(
            input_tokens=1, output_tokens=1, duration_seconds=0.1,
            system=tracker.SYSTEM_REASONING, provider="openai", model="gpt-4o-mini",
        )
        row = _rows(log_file)[1]
        assert _cell(row, "System") == tracker.SYSTEM_REASONING
        assert _cell(row, "Model") == "gpt-4o-mini"


class TestSeparateTotalsPerModel:
    """Each model gets its own colour and its own subtotal, so a call on an
    expensive model is never folded into the everyday model's total."""

    def _log(self, model, tokens):
        tracker.log_usage(
            input_tokens=tokens, output_tokens=tokens // 10, duration_seconds=1.0, model=model
        )

    def test_one_subtotal_per_model_then_the_grand_total(self, _isolated_files):
        _, log_file = _isolated_files
        self._log("gpt-4o-mini", 100)
        self._log("gpt-5.6-luna", 5000)
        self._log("gpt-4o-mini", 200)
        self._log("o3-mini", 50)

        rows = _rows(log_file)
        # Calls stay in the order they were made.
        assert [_cell(r, "Model") for r in rows[1:5]] == [
            "gpt-4o-mini", "gpt-5.6-luna", "gpt-4o-mini", "o3-mini",
        ]
        subtotals = rows[5:8]
        assert [_cell(r, "Model") for r in subtotals] == ["gpt-4o-mini", "gpt-5.6-luna", "o3-mini"]
        assert subtotals[1][0] == '="TOTAL gpt-5.6-luna calls: "&COUNTIF($F$2:$F$5,"gpt-5.6-luna")'
        letter = _letter("Input Tokens")
        assert _cell(subtotals[0], "Input Tokens") == f'=SUMIF($F$2:$F$5,"gpt-4o-mini",{letter}2:{letter}5)'
        assert rows[8][0] == '="TOTAL calls: "&COUNTA(A2:A5)'
        assert len(rows) == 9

    def test_a_model_s_rows_and_its_subtotal_share_one_colour(self, _isolated_files):
        _, log_file = _isolated_files
        self._log("gpt-4o-mini", 100)
        self._log("gpt-5.6-luna", 5000)
        self._log("gpt-4o-mini", 200)

        sheet = openpyxl.load_workbook(log_file).active
        colour = lambda row: sheet.cell(row=row, column=1).fill.start_color.rgb
        mini, luna = colour(2), colour(3)
        assert mini != luna
        assert colour(4) == mini          # the second gpt-4o-mini call
        assert colour(5) == mini          # gpt-4o-mini subtotal
        assert colour(6) == luna          # gpt-5.6-luna subtotal

    def test_the_reconciliation_still_counts_only_real_calls(self, _isolated_files):
        _, log_file = _isolated_files
        self._log("gpt-4o-mini", 100)
        self._log("gpt-5.6-luna", 5000)
        by_day = tracker._logged_totals_by_day(log_file)
        (day,) = by_day.values()
        assert day == {"calls": 2, "input_tokens": 5100, "output_tokens": 510}

    def test_format_recolours_an_old_single_total_log(self, _isolated_files):
        _, log_file = _isolated_files
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = tracker._SHEET_CALLS
        sheet.append(tracker._HEADERS)
        sheet.append(["2026-09-24T01:00:00+00:00", "fast", "fast_explanation", None, "openai", "gpt-4o-mini", 1.0, 100, 10, None])
        sheet.append(["2026-09-24T05:25:53+00:00", "fast", "fast_explanation", None, "openai", "gpt-5.6-luna", 6.4, 4868, 754, None])
        sheet.append(['="TOTAL calls: "&COUNTA(A2:A3)', None, None, None, None, None, "=SUM(G2:G3)", "=SUM(H2:H3)", "=SUM(I2:I3)", "=SUM(J2:J3)"])
        workbook.save(log_file)

        assert tracker.refresh_layout() is True
        rows = _rows(log_file)
        assert [_cell(r, "Model") for r in rows[3:5]] == ["gpt-4o-mini", "gpt-5.6-luna"]
        assert rows[-1][0] == '="TOTAL calls: "&COUNTA(A2:A3)'
        assert len(rows) == 6
>>>>>>> 8e557ed (Update Daily Report Agentic implementation)
